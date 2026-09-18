"""Рендеринг печатных форм из обновляемых шаблонов (п. 1.5, 7 ТЗ).

HTML-шаблоны лежат в каталоге settings.report_templates_dir и обновляются
отдельно от ядра (подписка Продукта 3 доставляет новые формы — достаточно
заменить файлы каталога, без пересборки образа).

Экспорт: HTML (печать/в PDF средствами браузера), XLSX, JSON
(машиночитаемая структура под будущие форматы ФНС).
"""

from __future__ import annotations

import io
from typing import Any

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook

from connector.config import settings


def _env(templates_dir: str | None = None) -> Environment:
    return Environment(
        loader=FileSystemLoader(templates_dir or settings.report_templates_dir),
        autoescape=select_autoescape(["html"]),
    )


def render_html(template_name: str, context: dict, templates_dir: str | None = None) -> str:
    return _env(templates_dir).get_template(template_name).render(**context)


def _autofit(ws) -> None:
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 70)


def _workbook_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def payment_act_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Акт по платежу"
    payment, imm = data["payment"], data["immutability"]
    rows: list[tuple[Any, Any]] = [
        ("Акт по операции с цифровой валютой", ""),
        ("Сформирован", data["generated_at"]),
        ("", ""),
        ("Направление", "поступление" if payment["direction"] == "in" else "выбытие"),
        ("Актив", payment["asset"]),
        ("Сумма", payment["amount"]),
        ("Отправитель", payment["from_address"]),
        ("Получатель", payment["to_address"]),
        ("Время блока", payment["block_time"]),
        ("Финализирована", payment["finalized_at"]),
        ("Подтверждений", payment["confirmations"]),
    ]
    if data["rate"]:
        rate = data["rate"]
        rows += [
            ("", ""),
            (f"Курс {payment['asset']}/{rate['contract_currency']}", rate["asset_to_contract"]),
            (f"Курс {rate['contract_currency']}/RUB", rate["contract_to_rub"]),
            ("Сумма, руб.", rate["amount_rub"]),
            ("Источник курса", rate["source"]),
        ]
    rows += [("", ""), ("Привязки (валютный контроль)", "")]
    for alloc in data["allocations"]:
        rows.append(
            (
                f"{alloc['counterparty'] or '—'} / контракт {alloc['contract_number'] or '—'}"
                f" (уч. № {alloc['contract_registration_number'] or '—'})",
                alloc["amount"],
            )
        )
    rows += [
        ("", ""),
        ("Журнал неизменяемости", ""),
        ("Сеть", imm["network"]),
        ("Хэш транзакции", imm["tx_hash"]),
        ("Блок", imm["block_number"]),
        ("Источник данных", imm["source"]),
        ("Получено", imm["received_at"]),
        ("SHA-256 сырого ответа ноды", imm["raw_response_sha256"]),
    ]
    for row in rows:
        ws.append(row)
    _autofit(ws)
    return _workbook_bytes(wb)


KIND_RU = {"receipt": "поступление", "disposal": "выбытие", "fee": "комиссия сети"}


def tax_register_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Налоговый регистр"
    ws.append(["Налоговый регистр операций с цифровой валютой"])
    ws.append([f"Период: {data['period']['from']} — {data['period']['to']}"])
    ws.append([])
    ws.append(["Дата (UTC)", "Операция", "Актив", "Количество",
               "Доход, ₽", "Расход, ₽", "Результат, ₽", "Хэш транзакции"])
    for row in data["rows"]:
        ws.append([
            row["date"], KIND_RU.get(row["kind"], row["kind"]), row["asset"],
            row["quantity"], row["income_rub"], row["cost_rub"], row["result_rub"],
            row["tx_hash"],
        ])
    ws.append([])
    ws.append(["Доходы, ₽", data["totals"]["income_rub"]])
    ws.append(["Расходы, ₽", data["totals"]["expense_rub"]])
    ws.append(["Налоговая база, ₽", data["totals"]["result_rub"]])
    _autofit(ws)
    return _workbook_bytes(wb)


def reconciliation_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Акт сверки"
    ws.append([f"Акт сверки / Reconciliation: {data['counterparty']['name']}"])
    ws.append([f"Период / Period: {data['period']['from']} — {data['period']['to']}"])
    ws.append([])
    ws.append(["Начислено / Invoiced"])
    ws.append(["Дата", "Инвойс", "Контракт", "Сумма", "Валюта"])
    for i in data["invoices"]:
        ws.append([i["date"], i["number"], i["contract_number"], i["amount"], i["currency"]])
    ws.append([])
    ws.append(["Оплачено / Paid"])
    ws.append(["Дата", "Актив", "Сумма", "В валюте контракта", "Контракт", "Хэш"])
    for p in data["payments"]:
        ws.append([p["date"], p["asset"], p["amount"],
                   p["contract_currency_amount"], p["contract_number"], p["tx_hash"]])
    ws.append([])
    ws.append(["Начислено / Invoiced", data["totals"]["invoiced"], data["totals"]["currency"]])
    ws.append(["Оплачено / Paid", data["totals"]["paid"], data["totals"]["currency"]])
    ws.append(["Сальдо / Balance", data["totals"]["balance"], data["totals"]["currency"]])
    _autofit(ws)
    return _workbook_bytes(wb)


def journal_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал операций"
    ws.append([f"Журнал операций по адресу {data['wallet']['address']}"])
    ws.append([f"Сеть: {data['wallet']['network']}",
               f"Период: {data['period']['from']} — {data['period']['to']}"])
    ws.append([])
    ws.append(
        [
            "Время блока", "Направление", "Актив", "Сумма", "Сумма, руб.",
            "Отправитель", "Получатель", "Контрагент", "Контракт (уч. №)",
            "Хэш транзакции", "SHA-256 ответа ноды",
        ]
    )
    for row in data["rows"]:
        alloc = row["allocations"][0] if row["allocations"] else {}
        contract = alloc.get("contract_number")
        reg = alloc.get("contract_registration_number")
        ws.append(
            [
                row["block_time"],
                "поступление" if row["direction"] == "in" else "выбытие",
                row["asset"],
                row["amount"],
                row["amount_rub"],
                row["from_address"],
                row["to_address"],
                alloc.get("counterparty"),
                f"{contract} ({reg})" if contract else None,
                row["immutability"]["tx_hash"],
                row["immutability"]["raw_response_sha256"],
            ]
        )
    ws.append([])
    ws.append(["Итого операций", data["totals"]["count"]])
    ws.append(["Поступило", data["totals"]["in"]])
    ws.append(["Выбыло", data["totals"]["out"]])
    _autofit(ws)
    return _workbook_bytes(wb)
