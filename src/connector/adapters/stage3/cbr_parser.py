"""Verified-parser boundary helpers for CBR registry evidence.

The JSON schema implemented here is an internal normalized interchange
format used for frozen fixtures and controlled imports. It is NOT asserted
to be the native publication format of the Bank of Russia.

A native CBR parser must be added only after the real publication format
has been captured and verified.
"""

from __future__ import annotations

import json
import re
from io import BytesIO
from datetime import datetime, timedelta, timezone
from typing import Any

from openpyxl import load_workbook

from connector.adapters.stage3.adapters import RegistryResult
from connector.adapters.stage3.cbr_registry import (
    CBRRawResponse,
    CBRRegistryRequest,
)


def _parse_dt(value: str | None) -> datetime | None:
    if not value:
        return None
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


class NormalizedCBRFixtureParser:
    """Parser for the connector's frozen normalized fixture format."""

    def parse(
        self,
        request: CBRRegistryRequest,
        response: CBRRawResponse,
    ) -> RegistryResult:
        try:
            payload: dict[str, Any] = json.loads(response.content)
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("CBR fixture is not valid JSON") from exc

        subject_ref = str(payload.get("subject_ref") or "")
        if not subject_ref:
            raise ValueError("CBR fixture missing subject_ref")

        if subject_ref != request.subject_ref:
            raise ValueError(
                "CBR fixture subject_ref does not match requested subject"
            )

        registry_type = str(payload.get("registry_type") or "")
        if registry_type != request.registry_type.value:
            raise ValueError(
                "CBR fixture registry_type does not match requested registry"
            )

        status = str(payload.get("status") or "")
        if not status:
            raise ValueError("CBR fixture missing status")

        entry_id = payload.get("registry_entry_id")
        admission_mode = payload.get("admission_mode")

        return RegistryResult(
            subject_ref=subject_ref,
            status=status,
            registry_type=registry_type,
            record_number=str(entry_id) if entry_id else None,
            registry_entry_id=str(entry_id) if entry_id else None,
            admission_mode=(
                str(admission_mode) if admission_mode else None
            ),
            valid_from=_parse_dt(payload.get("valid_from")),
            valid_to=_parse_dt(payload.get("valid_to")),
            application_id=(
                str(payload["application_id"])
                if payload.get("application_id")
                else None
            ),
            admission_date=_parse_dt(payload.get("admission_date")),
            legal_basis=(
                str(payload["legal_basis"])
                if payload.get("legal_basis")
                else None
            ),
            checked_at=response.checked_at,
            raw_payload=payload,
            checksum=response.checksum,
            source="CBR",
            source_url=response.source_url,
        )

MOSCOW_TZ = timezone(timedelta(hours=3))

_OIS_SHEET = "Реестр ОИС"
_OIS_TITLE_PREFIX = (
    "Реестр операторов информационных систем, "
    "в которых осуществляется выпуск цифровых финансовых активов"
)


def _clean_cell(value: Any) -> str | None:
    if value is None:
        return None

    text = str(value).strip()

    if not text or text == "-":
        return None

    return text


def _parse_cbr_date(value: Any) -> datetime | None:
    text = _clean_cell(value)
    if text is None:
        return None

    try:
        parsed = datetime.strptime(text, "%d.%m.%Y")
    except ValueError as exc:
        raise ValueError(
            f"Unexpected CBR registry date format: {text!r}"
        ) from exc

    return parsed.replace(tzinfo=MOSCOW_TZ)


def _extract_registry_as_of(value: Any) -> datetime:
    text = _clean_cell(value)
    if text is None:
        raise ValueError("CBR OIS registry is missing as-of date")

    match = re.search(r"(\d{2}\.\d{2}\.\d{4})", text)
    if match is None:
        raise ValueError(
            f"Cannot parse CBR OIS registry as-of date: {text!r}"
        )

    parsed = _parse_cbr_date(match.group(1))
    assert parsed is not None
    return parsed


def _ois_status(
    exclusion_decision: Any,
    exclusion_record: Any,
) -> str:
    """Derive only what the official XLSX itself proves.

    A recorded exclusion is final evidence of removal from the registry.
    A decision without the corresponding registry record is kept as a
    separate conservative state instead of being treated as active.
    """

    if _clean_cell(exclusion_record):
        return "excluded"

    if _clean_cell(exclusion_decision):
        return "exclusion_decided"

    return "active"


class CBROISXlsxParser:
    """Native parser for the official Bank of Russia OIS XLSX registry.

    Verified against the public workbook structure observed on 2026-10-05:
    sheet "Реестр ОИС", fourteen numbered columns, data starting at row 6.

    Subject lookup accepts either INN or OGRN because both identifiers are
    explicitly published in the workbook.

    The workbook does not publish a separate registry-entry number or
    admission mode; those fields therefore remain None rather than being
    invented by the connector.
    """

    def parse(
        self,
        request: CBRRegistryRequest,
        response: CBRRawResponse,
    ) -> RegistryResult:
        if (
            request.registry_type.value
            != "INFORMATION_SYSTEM_OPERATOR"
        ):
            raise ValueError(
                "CBROISXlsxParser can parse only "
                "INFORMATION_SYSTEM_OPERATOR registry"
            )

        try:
            workbook = load_workbook(
                BytesIO(response.content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise ValueError(
                "CBR OIS response is not a valid XLSX workbook"
            ) from exc

        if _OIS_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"CBR OIS workbook missing sheet {_OIS_SHEET!r}"
            )

        ws = workbook[_OIS_SHEET]

        title = _clean_cell(ws.cell(row=1, column=1).value)
        if not title or not title.startswith(_OIS_TITLE_PREFIX):
            raise ValueError(
                "CBR OIS workbook title does not match verified format"
            )

        registry_as_of = _extract_registry_as_of(
            ws.cell(row=2, column=1).value
        )

        # Fail closed if the official workbook changes its structural
        # contract. Do not silently parse shifted columns.
        if _clean_cell(ws.cell(row=3, column=3).value) != "ОГРН":
            raise ValueError("CBR OIS workbook OGRN column moved")

        if _clean_cell(ws.cell(row=3, column=4).value) != "ИНН":
            raise ValueError("CBR OIS workbook INN column moved")

        column_numbers = [
            ws.cell(row=5, column=i).value
            for i in range(1, 15)
        ]
        if column_numbers != list(range(1, 15)):
            raise ValueError(
                "CBR OIS workbook column numbering changed"
            )

        matches: list[tuple[Any, ...]] = []

        for row in ws.iter_rows(
            min_row=6,
            max_col=14,
            values_only=True,
        ):
            if not any(_clean_cell(value) for value in row):
                continue

            ogrn = _clean_cell(row[2])
            inn = _clean_cell(row[3])

            if request.subject_ref in {inn, ogrn}:
                matches.append(row)

        if len(matches) > 1:
            raise ValueError(
                "CBR OIS registry contains multiple rows for "
                f"subject {request.subject_ref}"
            )

        if not matches:
            return RegistryResult(
                subject_ref=request.subject_ref,
                status="not_found",
                registry_type=request.registry_type.value,
                record_number=None,
                registry_entry_id=None,
                checked_at=response.checked_at,
                raw_payload={
                    "sheet": _OIS_SHEET,
                    "registry_as_of": registry_as_of.isoformat(),
                    "matched": False,
                    "subject_ref": request.subject_ref,
                },
                checksum=response.checksum,
                source="CBR",
                source_url=response.source_url,
            )

        row = matches[0]

        inclusion_date = _parse_cbr_date(row[0])
        exclusion_decision_date = _parse_cbr_date(row[12])
        exclusion_record_date = _parse_cbr_date(row[13])

        status = _ois_status(
            row[12],
            row[13],
        )

        raw_row = {
            "included_at": _clean_cell(row[0]),
            "name": _clean_cell(row[1]),
            "ogrn": _clean_cell(row[2]),
            "inn": _clean_cell(row[3]),
            "egrul_address": _clean_cell(row[4]),
            "website": _clean_cell(row[5]),
            "phone": _clean_cell(row[6]),
            "email": _clean_cell(row[7]),
            "platform_registry_date": _clean_cell(row[8]),
            "platform_right_acquired_date": _clean_cell(row[9]),
            "platform_exclusion_date": _clean_cell(row[10]),
            "platform_right_lost_date": _clean_cell(row[11]),
            "exclusion_decision_date": _clean_cell(row[12]),
            "exclusion_record_date": _clean_cell(row[13]),
        }

        return RegistryResult(
            subject_ref=request.subject_ref,
            status=status,
            registry_type=request.registry_type.value,

            # The official workbook contains no independent registry
            # entry number. Do not substitute INN/OGRN as a fake one.
            record_number=None,
            registry_entry_id=None,

            # The OIS XLSX does not publish general/simplified/etc.
            admission_mode=None,

            valid_from=inclusion_date,
            valid_to=exclusion_record_date,

            application_id=None,
            admission_date=inclusion_date,
            legal_basis=None,

            checked_at=response.checked_at,
            raw_payload={
                "sheet": _OIS_SHEET,
                "registry_as_of": registry_as_of.isoformat(),
                "matched": True,
                "record": raw_row,
                "exclusion_decision_at": (
                    exclusion_decision_date.isoformat()
                    if exclusion_decision_date
                    else None
                ),
            },
            checksum=response.checksum,
            source="CBR",
            source_url=response.source_url,
        )
