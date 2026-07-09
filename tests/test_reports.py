"""Тесты печатных форм: сборка данных, журнал неизменяемости, рендеринг (п. 7 ТЗ)."""

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from connector.models import (
    Asset,
    Base,
    Contract,
    Counterparty,
    CounterpartyAddress,
    Direction,
    Invoice,
    Network,
    Transaction,
    TxStatus,
    Wallet,
)
from connector.pipeline import TransactionPipeline
from connector.rates.service import RateService
from connector.reports.render import journal_xlsx, payment_act_xlsx, render_html
from connector.reports.service import journal_data, payment_act_data, raw_response_sha256

from tests.test_pipeline import FakeRateSource  # переиспользуем фейковый источник

TEMPLATES_DIR = str(Path(__file__).parent.parent / "templates" / "reports")


def dt(day: int, hour: int = 12) -> datetime:
    return datetime(2026, 6, day, hour, tzinfo=timezone.utc)


@pytest.fixture
async def session():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as session:
        yield session
    await engine.dispose()


@pytest.fixture
async def processed_tx(session):
    """Финализированная входящая транзакция, проведённая через конвейер."""
    network = Network(code="tron", name="TRON", finality_depth=19)
    session.add(network)
    await session.flush()
    asset = Asset(network_id=network.id, symbol="USDT", contract_address="TUsdt", decimals=6)
    wallet = Wallet(network_id=network.id, address="TMyWallet", label="Основной")
    counterparty = Counterparty(name="Shanghai Trade Co.")
    session.add_all([asset, wallet, counterparty])
    await session.flush()
    session.add(
        CounterpartyAddress(
            counterparty_id=counterparty.id, network_id=network.id, address="TPartner"
        )
    )
    contract = Contract(
        counterparty_id=counterparty.id,
        number="ВЭД-1",
        registration_number="21060001/1481/0000/9/1",
        currency="USD",
    )
    session.add(contract)
    await session.flush()
    session.add(
        Invoice(
            contract_id=contract.id,
            number="INV-1",
            amount=Decimal(1000),
            currency="USD",
            due_from=dt(1),
            due_to=dt(28),
        )
    )
    tx = Transaction(
        network_id=network.id,
        asset_id=asset.id,
        wallet_id=wallet.id,
        tx_hash="abc123",
        log_index=0,
        block_number=1010,
        block_time=dt(10),
        direction=Direction.IN,
        from_address="TPartner",
        to_address="TMyWallet",
        amount=Decimal(1000),
        status=TxStatus.FINAL,
        confirmations=25,
        finalized_at=dt(10, 13),
        raw_response={"transaction_id": "abc123", "value": "1000000000"},
        source="tron-1",
    )
    session.add(tx)
    await session.flush()
    pipeline = TransactionPipeline(session, RateService(primary=FakeRateSource()))
    await pipeline.process_network(network)
    return {"tx": tx, "wallet": wallet, "network": network}


def test_raw_hash_is_canonical():
    # Порядок ключей не влияет на хэш — сверка при проверках воспроизводима.
    assert raw_response_sha256({"a": 1, "b": 2}) == raw_response_sha256({"b": 2, "a": 1})
    assert raw_response_sha256({"a": 1}) != raw_response_sha256({"a": 2})


async def test_payment_act_data_includes_control_and_immutability(session, processed_tx):
    data = await payment_act_data(session, processed_tx["tx"].id)
    assert data["payment"]["amount"] == "1000"
    assert Decimal(data["rate"]["amount_rub"]) == Decimal(80000)
    [alloc] = data["allocations"]
    assert alloc["counterparty"] == "Shanghai Trade Co."
    assert alloc["contract_registration_number"] == "21060001/1481/0000/9/1"
    imm = data["immutability"]
    assert imm["tx_hash"] == "abc123"
    assert imm["raw_response_sha256"] == raw_response_sha256(
        {"transaction_id": "abc123", "value": "1000000000"}
    )


async def test_payment_act_renders_html_and_xlsx(session, processed_tx):
    data = await payment_act_data(session, processed_tx["tx"].id)
    html = render_html("payment_act.html", data, templates_dir=TEMPLATES_DIR)
    assert "АКТ" in html
    assert "abc123" in html
    assert "21060001/1481/0000/9/1" in html

    wb = load_workbook(BytesIO(payment_act_xlsx(data)))
    cells = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value]
    assert any("abc123" in c for c in cells)


async def test_journal_covers_period_with_immutability(session, processed_tx):
    data = await journal_data(session, processed_tx["wallet"].id, dt(1), dt(30))
    assert data["totals"]["count"] == 1
    [row] = data["rows"]
    assert row["immutability"]["tx_hash"] == "abc123"
    assert Decimal(row["amount_rub"]) == Decimal(80000)

    # За пустой период — пустой журнал, а не ошибка.
    empty = await journal_data(session, processed_tx["wallet"].id, dt(20), dt(30))
    assert empty["rows"] == []


async def test_journal_renders_html_and_xlsx(session, processed_tx):
    data = await journal_data(session, processed_tx["wallet"].id, dt(1), dt(30))
    html = render_html("operations_journal.html", data, templates_dir=TEMPLATES_DIR)
    assert "ЖУРНАЛ ОПЕРАЦИЙ" in html
    assert "TMyWallet" in html

    wb = load_workbook(BytesIO(journal_xlsx(data)))
    cells = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value]
    assert any("abc123" in c for c in cells)


async def test_missing_entities_return_none(session):
    assert await payment_act_data(session, 999) is None
    assert await journal_data(session, 999, dt(1), dt(30)) is None
